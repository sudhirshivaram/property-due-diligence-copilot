"""Extract ordered paragraphs and tables."""

from io import BytesIO
from hashlib import sha256
from zipfile import ZipFile
from collections import Counter
import xml.etree.ElementTree as ET
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.enum.text import WD_BREAK
from .primitives import (
    all_paragraph_records,
    base_record,
    preview_record,
    property_xml,
    show_full_record,
    show_rows,
    table_paragraphs,
)


class ExtractionSteps:
    """Extraction steps; state belongs to the workflow instance."""

    def source_style(self, xml, kind):
        style_path = "w:pPr/w:pStyle" if kind == "paragraph" else "w:tblPr/w:tblStyle"
        node = xml.find(style_path, self.NS)
        explicit = node.get(self.W + "val") if node is not None else None
        resolved = explicit if explicit is not None else self.default_styles.get(kind)
        return {
            "explicit_id": explicit,
            "resolved_id": resolved,
            "name": self.style_catalog.get(resolved),
        }

    def run_events(self, run_xml):
        events = []
        for child in run_xml:
            if child.tag == self.W + "rPr":
                continue
            if child.tag == self.W + "t":
                events.append({"type": "text", "text": child.text or ""})
            elif child.tag == self.W + "tab":
                events.append({"type": "tab"})
            elif child.tag in {self.W + "br", self.W + "cr"}:
                events.append(
                    {
                        "type": "break",
                        "xml_tag": child.tag.rsplit("}", 1)[-1],
                        "kind": child.get(self.W + "type", "textWrapping"),
                        "clear": child.get(self.W + "clear"),
                    }
                )
            elif child.tag == self.W + "fldChar":
                events.append(
                    {
                        "type": "field_char",
                        "kind": child.get(self.W + "fldCharType"),
                        "xml": property_xml(child),
                    }
                )
            elif child.tag == self.W + "instrText":
                events.append({"type": "field_instruction", "text": child.text or ""})
            else:
                events.append(
                    {
                        "type": "unhandled_xml",
                        "tag": child.tag,
                        "xml": property_xml(child),
                    }
                )
        return events

    def step_4_extract_ordered_structural_blocks(self):
        """Step 4 — Extract ordered structural blocks."""
        self.style_catalog = {s.style_id: s.name for s in self.parser_document.styles}
        self.default_styles = {
            s.get(self.W + "type"): s.get(self.W + "styleId")
            for s in self.styles_xml.findall("w:style", self.NS)
            if s.get(self.W + "default") == "1"
        }

    def paragraph_record(self, paragraph, xml, position, part, path, source_hash):
        unexpected = [e.tag for e in xml if e.tag not in {self.W + "pPr", self.W + "r"}]
        if unexpected:
            raise ValueError(
                f"Review unsupported paragraph content at {part}:{path}: {unexpected}"
            )
        record = base_record("paragraph", position, part, path, source_hash)
        record.update(
            raw_text=paragraph.text,
            style=self.source_style(xml, "paragraph"),
            paragraph_properties_xml=property_xml(xml.find("w:pPr", self.NS)),
            runs=[],
            break_markers=[],
        )
        xml_runs = xml.findall("w:r", self.NS)
        for run_index, (run, run_xml) in enumerate(
            zip(paragraph.runs, xml_runs, strict=True)
        ):
            underline = run.underline
            if underline is not None and (not isinstance(underline, bool)):
                underline = underline.name
            formatting = {
                "bold": run.bold,
                "italic": run.italic,
                "underline": underline,
                "strike": run.font.strike,
                "superscript": run.font.superscript,
                "subscript": run.font.subscript,
                "font_name": run.font.name,
                "font_size_pt": run.font.size.pt if run.font.size is not None else None,
            }
            events = self.run_events(run_xml)
            record["runs"].append(
                {
                    "position": run_index,
                    "raw_text": run.text,
                    "formatting": formatting,
                    "properties_xml": property_xml(run_xml.find("w:rPr", self.NS)),
                    "events": events,
                }
            )
            for event_index, event in enumerate(events):
                if event["type"] == "break":
                    record["break_markers"].append(
                        {
                            "run_index": run_index,
                            "event_index": event_index,
                            "kind": event["kind"],
                            "xml_tag": event["xml_tag"],
                            "clear": event["clear"],
                        }
                    )
        return record

    def table_record(self, table, xml, position, part, path, source_hash):
        if list(xml.iter(self.W + "gridSpan")) or list(xml.iter(self.W + "vMerge")):
            raise ValueError(f"Merged cells require review at {part}:{path}")
        if len(list(xml.iter(self.W + "tbl"))) != 1:
            raise ValueError(f"Nested tables require review at {part}:{path}")
        if any(
            (
                e.tag not in {self.W + "tblPr", self.W + "tblGrid", self.W + "tr"}
                for e in xml
            )
        ):
            raise ValueError(f"Unexpected table content at {part}:{path}")
        record = base_record("table", position, part, path, source_hash)
        record.update(style=self.source_style(xml, "table"), rows=[])
        for row_index, (row, row_xml) in enumerate(
            zip(table.rows, xml.findall("w:tr", self.NS), strict=True)
        ):
            if any((e.tag not in {self.W + "trPr", self.W + "tc"} for e in row_xml)):
                raise ValueError(
                    f"Unexpected row content at {part}:{path}/rows/{row_index}"
                )
            if row.grid_cols_before or row.grid_cols_after:
                raise ValueError(
                    f"Omitted grid cells require review at {part}:{path}/rows/{row_index}"
                )
            row_record = {"row_index": row_index, "cells": []}
            for column_index, (cell, cell_xml) in enumerate(
                zip(row.cells, row_xml.findall("w:tc", self.NS), strict=True)
            ):
                cell_path = f"{path}/rows/{row_index}/cells/{column_index}"
                if any(
                    (e.tag not in {self.W + "tcPr", self.W + "p"} for e in cell_xml)
                ):
                    raise ValueError(f"Unexpected cell content at {part}:{cell_path}")
                cell_paragraphs = []
                api_paragraphs = iter(cell.paragraphs)
                for child_index, child in enumerate(cell_xml):
                    if child.tag == self.W + "p":
                        cell_paragraphs.append(
                            self.paragraph_record(
                                next(api_paragraphs),
                                child,
                                child_index,
                                part,
                                f"{cell_path}/children/{child_index}",
                                source_hash,
                            )
                        )
                if next(api_paragraphs, None) is not None:
                    raise AssertionError(f"Unaccounted cell paragraph at {cell_path}")
                row_record["cells"].append(
                    {"column_index": column_index, "paragraphs": cell_paragraphs}
                )
            record["rows"].append(row_record)
        return record

    def extract_container(self, items, xml_container, part, path, source_hash):
        """Keep source order; XML positions include non-content layout elements."""
        items = iter(items)
        blocks, layout = ([], [])
        for position, xml in enumerate(xml_container):
            location = f"{path}/children/{position}"
            if xml.tag == self.W + "sectPr":
                marker = base_record(
                    "section_properties", position, part, location, source_hash
                )
                marker["xml"] = property_xml(xml)
                layout.append(marker)
                continue
            if xml.tag not in {self.W + "p", self.W + "tbl"}:
                raise ValueError(
                    f"Unsupported container element at {part}:{location}: {xml.tag}"
                )
            item = next(items, None)
            if xml.tag == self.W + "p" and isinstance(item, Paragraph):
                blocks.append(
                    self.paragraph_record(
                        item, xml, position, part, location, source_hash
                    )
                )
            elif xml.tag == self.W + "tbl" and isinstance(item, Table):
                blocks.append(
                    self.table_record(item, xml, position, part, location, source_hash)
                )
            else:
                raise AssertionError(f"Parser/XML order mismatch at {part}:{location}")
        if next(items, None) is not None:
            raise AssertionError(f"Unaccounted parser content in {part}")
        return (blocks, layout)

    def step_4c_tables_remain_in_place_with_nested_cell_paragraphs(self):
        """4c. Tables remain in place, with nested cell paragraphs."""
        self.ordered_blocks, self.layout_markers = self.extract_container(
            self.parser_document.iter_inner_content(),
            self.body,
            "word/document.xml",
            "/body",
            self.manifest["sha256"],
        )
        print("Body records:", dict(Counter((b["type"] for b in self.ordered_blocks))))
        print("Separate layout markers:", len(self.layout_markers))

    def step_4d_footer_content_stays_outside_the_body(self):
        """4d. Footer content stays outside the body."""
        self.parts_by_name = {
            str(part.partname).lstrip("/"): part
            for part in self.parser_document.part.package.parts
        }
        self.auxiliary_content = []
        for self.part_name, self.xml in self.auxiliary_xml.items():
            self.part = self.parts_by_name[self.part_name]
            self.items = (
                Paragraph(child, self.part)
                if child.tag == self.W + "p"
                else Table(child, self.part)
                for child in self.part.element
                if child.tag in {self.W + "p", self.W + "tbl"}
            )
            self.records, self.layout = self.extract_container(
                self.items,
                self.xml,
                self.part_name,
                "/" + self.xml.tag.rsplit("}", 1)[-1],
                self.manifest["sha256"],
            )
            self.auxiliary_content.append(
                {
                    "part": self.part_name,
                    "blocks": self.records,
                    "layout_markers": self.layout,
                }
            )
        show_rows(
            [
                {"part": a["part"], "paragraph/table records": len(a["blocks"])}
                for a in self.auxiliary_content
            ]
        )

    def step_4e_inspect_the_records(self):
        """4e. Inspect the records."""
        show_rows(
            [
                preview_record(b)
                for b in self.ordered_blocks
                if self.first_table_index - 3
                <= b["position"]
                <= self.first_table_index + 3
            ]
        )
        self.first_table_record = next(
            (b for b in self.ordered_blocks if b["type"] == "table")
        )
        show_rows(
            [
                {
                    "row": row["row_index"],
                    "column": cell["column_index"],
                    "paragraph position": p["position"],
                    "raw text": p["raw_text"],
                }
                for row in self.first_table_record["rows"]
                for cell in row["cells"]
                for p in cell["paragraphs"]
            ]
        )
        self.SELECTED_BODY_POSITION = self.first_table_index - 2
        self.selected = next(
            (
                b
                for b in self.ordered_blocks
                if b["position"] == self.SELECTED_BODY_POSITION
            )
        )
        show_full_record(self.selected, "Full selected paragraph record")
        show_full_record(
            self.first_table_record,
            "Full first table record (nested cells and paragraphs)",
        )
        self.page_break_paragraph = next(
            (
                b
                for b in self.ordered_blocks
                if b["type"] == "paragraph"
                and b["raw_text"] == ""
                and b["break_markers"]
            )
        )
        show_full_record(
            self.page_break_paragraph, "Empty visible text, preserved page-break event"
        )
        for self.auxiliary in self.auxiliary_content:
            show_full_record(
                self.auxiliary, "Separate auxiliary content: " + self.auxiliary["part"]
            )

    def xml_visible_text(self, xml):
        pieces = []
        for element in xml.iter():
            if element.tag == self.W + "t":
                pieces.append(element.text or "")
            elif element.tag == self.W + "tab":
                pieces.append("\t")
            elif element.tag == self.W + "cr" or (
                element.tag == self.W + "br"
                and element.get(self.W + "type", "textWrapping") == "textWrapping"
            ):
                pieces.append("\n")
        return "".join(pieces)

    def step_4f_check_step_4_preservation(self):
        """4f. Check Step 4 preservation."""
        self.body_paragraph_records = all_paragraph_records(self.ordered_blocks)
        self.xml_paragraphs = list(self.body.iter(self.W + "p"))
        for self.record, self.xml in zip(
            self.body_paragraph_records, self.xml_paragraphs, strict=True
        ):
            assert self.record["raw_text"] == self.xml_visible_text(self.xml), (
                self.record["source"]
            )
            assert self.record["raw_text"] == "".join(
                (r["raw_text"] for r in self.record["runs"])
            ), self.record["source"]
            assert [
                e["text"]
                for r in self.record["runs"]
                for e in r["events"]
                if e["type"] == "text"
            ] == [e.text or "" for e in self.xml.iter(self.W + "t")], self.record[
                "source"
            ]
            assert [
                (m["xml_tag"], m["kind"], m["clear"])
                for m in self.record["break_markers"]
            ] == [
                (
                    e.tag.rsplit("}", 1)[-1],
                    e.get(self.W + "type", "textWrapping"),
                    e.get(self.W + "clear"),
                )
                for e in self.xml.iter()
                if e.tag in {self.W + "br", self.W + "cr"}
            ], self.record["source"]
        assert [b["position"] for b in self.ordered_blocks] == [
            i
            for i, e in enumerate(self.body)
            if e.tag in {self.W + "p", self.W + "tbl"}
        ]
        assert [b["type"] for b in self.ordered_blocks] == [
            "paragraph" if e.tag == self.W + "p" else "table"
            for e in self.body
            if e.tag in {self.W + "p", self.W + "tbl"}
        ]
        assert len(self.ordered_blocks) + len(self.layout_markers) == len(self.body)
        self.all_records = (
            list(self.ordered_blocks)
            + [
                p
                for b in self.ordered_blocks
                if b["type"] == "table"
                for p in table_paragraphs(b)
            ]
            + self.layout_markers
        )
        for self.a in self.auxiliary_content:
            self.all_records.extend(self.a["blocks"])
            self.all_records.extend(
                (
                    p
                    for b in self.a["blocks"]
                    if b["type"] == "table"
                    for p in table_paragraphs(b)
                )
            )
            self.all_records.extend(self.a["layout_markers"])
        assert len({r["id"] for r in self.all_records}) == len(self.all_records)
        assert all(
            (
                r["id"]
                == base_record(
                    r["type"],
                    r["position"],
                    r["source"]["part"],
                    r["source"]["path"],
                    self.manifest["sha256"],
                )["id"]
                for r in self.all_records
            )
        )
        self.all_paragraphs = self.body_paragraph_records + [
            p
            for a in self.auxiliary_content
            for p in all_paragraph_records(a["blocks"])
        ]
        self.unhandled = [
            e
            for p in self.all_paragraphs
            for r in p["runs"]
            for e in r["events"]
            if e["type"] == "unhandled_xml"
        ]
        assert not self.unhandled, (
            "Unhandled run XML was retained; review before calling extraction complete"
        )
        for self.a in self.auxiliary_content:
            self.aux_paragraphs = all_paragraph_records(self.a["blocks"])
            for self.record, self.xml in zip(
                self.aux_paragraphs,
                self.auxiliary_xml[self.a["part"]].iter(self.W + "p"),
                strict=True,
            ):
                assert self.record["raw_text"] == self.xml_visible_text(self.xml)
                assert [
                    e["text"]
                    for r in self.record["runs"]
                    for e in r["events"]
                    if e["type"] == "field_instruction"
                ] == [e.text or "" for e in self.xml.iter(self.W + "instrText")]
                assert [
                    e["kind"]
                    for r in self.record["runs"]
                    for e in r["events"]
                    if e["type"] == "field_char"
                ] == [
                    e.get(self.W + "fldCharType")
                    for e in self.xml.iter(self.W + "fldChar")
                ]
        self.summary = {
            "body paragraph records": sum(
                (b["type"] == "paragraph" for b in self.ordered_blocks)
            ),
            "body table records": sum(
                (b["type"] == "table" for b in self.ordered_blocks)
            ),
            "table rows": sum(
                (len(b["rows"]) for b in self.ordered_blocks if b["type"] == "table")
            ),
            "table cells": sum(
                (
                    len(r["cells"])
                    for b in self.ordered_blocks
                    if b["type"] == "table"
                    for r in b["rows"]
                )
            ),
            "all body/cell paragraph records": len(self.body_paragraph_records),
            "page breaks": sum(
                (
                    m["kind"] == "page"
                    for p in self.body_paragraph_records
                    for m in p["break_markers"]
                )
            ),
            "empty visible body paragraphs": sum(
                (
                    b["raw_text"] == ""
                    for b in self.ordered_blocks
                    if b["type"] == "paragraph"
                )
            ),
        }
        assert list(self.summary.values())[:6] == [
            self.counts[k]
            for k in (
                "top-level paragraphs",
                "top-level tables",
                "table rows",
                "table cells",
                "all body/table paragraphs",
                "explicit page breaks",
            )
        ]
        show_rows([{"measure": k, "preserved": v} for k, v in self.summary.items()])
        print(
            "PASS: order, source positions, text fragments, whitespace, breaks, IDs, and auxiliary field events"
        )

    def step_4f_check_step_4_preservation_2(self):
        """4f. Check Step 4 preservation."""
        self.fixture = Document()
        self.p = self.fixture.add_paragraph()
        self.p.add_run("  Начало\t ").bold = True
        self.r = self.p.add_run("before")
        self.r.italic = True
        self.r.add_break(WD_BREAK.PAGE)
        self.r.add_text("after")
        self.r.add_break(WD_BREAK.LINE)
        self.r.add_text("end  ")
        self.t = self.fixture.add_table(rows=1, cols=2)
        self.t.cell(0, 0).text = " cell 0 "
        self.t.cell(0, 0).add_paragraph("second paragraph")
        self.t.cell(0, 1).text = "Български"
        self.fixture.add_paragraph("")
        self.fixture.add_paragraph("tail")
        self.buffer = BytesIO()
        self.fixture.save(self.buffer)
        self.fixture_bytes = self.buffer.getvalue()
        self.fixture_hash = sha256(self.fixture_bytes).hexdigest()
        with ZipFile(BytesIO(self.fixture_bytes)) as self.z:
            self.fixture_body = ET.fromstring(self.z.read("word/document.xml")).find(
                "w:body", self.NS
            )
        self.fixture_parser = Document(BytesIO(self.fixture_bytes))
        self.fixture_blocks, self._ = self.extract_container(
            self.fixture_parser.iter_inner_content(),
            self.fixture_body,
            "word/document.xml",
            "/body",
            self.fixture_hash,
        )
        assert [b["type"] for b in self.fixture_blocks] == [
            "paragraph",
            "table",
            "paragraph",
            "paragraph",
        ]
        assert self.fixture_blocks[0]["raw_text"] == "  Начало\t beforeafter\nend  "
        assert self.fixture_blocks[0]["runs"][0]["formatting"]["bold"] is True
        assert self.fixture_blocks[0]["runs"][1]["formatting"]["italic"] is True
        assert self.fixture_blocks[0]["runs"][1]["formatting"]["bold"] is None
        assert [e["type"] for e in self.fixture_blocks[0]["runs"][1]["events"]] == [
            "text",
            "break",
            "text",
            "break",
            "text",
        ]
        assert [m["kind"] for m in self.fixture_blocks[0]["break_markers"]] == [
            "page",
            "textWrapping",
        ]
        assert [
            p["raw_text"]
            for p in self.fixture_blocks[1]["rows"][0]["cells"][0]["paragraphs"]
        ] == [" cell 0 ", "second paragraph"]
        assert self.fixture_blocks[2]["raw_text"] == ""
        assert self.fixture_blocks[3]["raw_text"] == "tail"
        self.fixture_again, self._ = self.extract_container(
            self.fixture_parser.iter_inner_content(),
            self.fixture_body,
            "word/document.xml",
            "/body",
            self.fixture_hash,
        )
        assert self.fixture_again == self.fixture_blocks
        show_rows([preview_record(b) for b in self.fixture_blocks])
        print(
            "PASS: in-memory fixture — interleaving, multiple cell paragraphs, exact whitespace, formatting, inline breaks, empty paragraph, repeatability"
        )
